"""Oyunun tüm eşya kütüphanesini tek bir Excel dosyasına yazar: meta/TD2_Kutuphane.xlsx

Kaynaklar (birleştirilir, sonra Excel'e yazılır):
  - meta/game_data.json                 programın kullandığı kütüphane
  - meta/sources/season_y8.json         son sezonlar (Rise Up, Into the Dark, Red Horizon) — Ubisoft sezon notları
  - meta/sources/wiki_exotics.json      exotic yuva/tür ve talent bilgisi (The Division Wiki, CC BY-SA)
  - meta/sources/mxswat/*.csv           talent açıklamaları, modlar, silah değerleri (mx-division-builds, CC BY-NC-SA 4.0)

Excel'deki bilgiler programa scripts/import_library_excel.py ile işlenir (Excel'de yapılan düzeltmeler de).

    python scripts/library_excel.py [çıktı.xlsx]
"""
import csv
import json
import re
import sys
from collections import defaultdict
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))
import gamedata  # noqa: E402

N = gamedata.nrm
SLOTS = ["Mask", "Backpack", "Chest", "Gloves", "Holster", "Kneepads"]
SLOT_TR = {"Mask": "Maske", "Backpack": "Sırt çantası", "Chest": "Göğüslük", "Gloves": "Eldiven", "Holster": "Holster",
           "Kneepads": "Dizlik", "Weapon": "Silah", "Sidearm": "Tabanca"}
QUAL = {"high-end": "High-End", "gear set": "Gear Set", "named": "Named", "exotic": "Exotic"}
WTYPES = {"assault rifle": "Assault Rifle", "rifle": "Rifle", "marksman rifle": "Marksman Rifle", "submachine gun": "Submachine Gun",
          "smg": "Submachine Gun", "light machine gun": "Light Machine Gun", "lmg": "Light Machine Gun", "shotgun": "Shotgun",
          "pistol": "Pistol", "light machinegun": "Light Machine Gun"}
GEAR_WORDS = {"mask": "Mask", "backpack": "Backpack", "body armor": "Chest", "chest": "Chest", "vest": "Chest", "gloves": "Gloves",
              "holster": "Holster", "knee pads": "Kneepads", "kneepads": "Kneepads"}


def load():
    d = json.loads((ROOT / "meta" / "game_data.json").read_text(encoding="utf-8"))
    season = json.loads((ROOT / "meta" / "sources" / "season_y8.json").read_text(encoding="utf-8"))
    wiki = json.loads((ROOT / "meta" / "sources" / "wiki_exotics.json").read_text(encoding="utf-8"))
    mx = {}
    for f in (ROOT / "meta" / "sources" / "mxswat").glob("*.csv"):
        mx[f.stem] = list(csv.DictReader(f.open(encoding="utf-8")))
    src = ROOT / "meta" / "sources"
    mx["_wiki_talents"] = json.loads((src / "wiki_talents.json").read_text(encoding="utf-8")) if (src / "wiki_talents.json").exists() else {}
    corr = json.loads((src / "corrections.json").read_text(encoding="utf-8")) if (src / "corrections.json").exists() else {}
    apply_corrections(d, corr)
    mx["_wiki_talents"] = {**mx["_wiki_talents"], **corr.get("talents", {})}
    mx["_talent_desc"] = json.loads((src / "wiki_talent_desc.json").read_text(encoding="utf-8")) if (src / "wiki_talent_desc.json").exists() else {}
    return d, season, wiki, mx


def apply_corrections(d, corr):
    """TD1 eşyalarını çıkar, yanlış kaliteleri düzelt (içe aktarıcı da aynısını yapar)."""
    gone = {N(x) for x in corr.get("remove_td1", [])}
    noexo = gone | {N(x) for x in corr.get("remove_from_exotics", [])}
    d["exotics"] = [x for x in d["exotics"] if N(x) not in noexo]
    d["named"] = [x for x in d["named"] if N(x) not in gone and N(x) not in {N(e) for e in d["exotics"]}]
    d["items"] = {k: v for k, v in d["items"].items() if N(k) not in gone}
    d["weapon_index"] = {k: v for k, v in d["weapon_index"].items() if N(k) not in gone}
    d["weapons"] = [x for x in d.get("weapons", []) if N(x) not in gone]
    for k, q in corr.get("weapon_quality", {}).items():
        for wk in d["weapon_index"]:
            if N(wk) == N(k): d["weapon_index"][wk]["quality"] = q
    for k, weapon_type in corr.get("weapon_type", {}).items():
        for wk in d["weapon_index"]:
            if N(wk) == N(k): d["weapon_index"][wk]["type"] = weapon_type
    for a, b in corr.get("exotic_alias", {}).items():
        if N(b) not in {N(x) for x in d["exotics"]} and d["weapon_index"].get(b, {}).get("quality") == "exotic":
            d["exotics"].append(b)
    return d


def wiki_talent(rec):
    t = (rec or {}).get("talent") or ""
    desc = (rec or {}).get("desc") or ""
    # wiki tablo işaretleri: "{| class=... |- !Icon !Description |- |left|thumb|100px |"
    desc = re.sub(r"\{\|[^|]*?(?=\|-)", " ", desc)
    desc = re.sub(r"\|-|!Icon|!Description|\|\}|\b(left|right)\|thumb\|\d+px\s*\|?", " ", desc)
    desc = re.sub(r"\(Tom Clancy's The Division( 2)?\)", "", desc)
    desc = re.sub(r"\s+", " ", desc).strip(" |")
    if not t or t.lower() in ("talent", "icon"):
        m = re.search(r"Talent\s*[-–:]\s*([^*•\n]{2,40}?)\s*(?:\*|•|$)", desc) or re.match(r"\s*([A-Z][\w'’!.\- ]{2,30}?)\s*(?:•|\*)", desc)
        t = m.group(1).strip() if m else ""
    t = t.strip(" :-|")
    if t:
        desc = re.sub(r"^\s*(Talent\s*[-–:]\s*)?" + re.escape(t) + r"\s*:?", "", desc).strip(" *:-|")
    desc = re.sub(r"\s*\*+\s*", " • ", desc).strip(" •")
    return t, desc[:300]


def wiki_kind(rec):
    ty = N((rec or {}).get("type") or "")
    if not ty: return None, None
    for w, s in GEAR_WORDS.items():
        if w in ty: return "gear", s
    for w, s in WTYPES.items():
        if ty.endswith(w) or ty == w: return "weapon", s
    return None, None


def build(d, season, wiki, mx):
    """Tüm tabloları satır listeleri olarak döndürür (Excel'e ve içe aktarıcıya aynı veri)."""
    items = {k: dict(v) for k, v in d["items"].items()}
    windex = {k: dict(v) for k, v in d["weapon_index"].items()}
    season_by = {N(x["name"]): x for x in season["items"]}
    mx_gear = {N(r["Item Name"]): r for r in mx.get("Gear", [])}
    mx_weap = {N(r["Name"]): r for r in mx.get("Weapons", [])}

    # sezon eşyaları kütüphaneye
    for x in season["items"]:
        if x["kind"] == "gear":
            rec = items.setdefault(next((k for k in items if N(k) == N(x["name"])), x["name"]), {})
            rec.update({"slot": x["slot"], "quality": x["category"].lower()})
            if x.get("brand"): rec["brand"] = x["brand"]
        else:
            key = next((k for k in windex if N(k) == N(x["name"])), x["name"])
            rec = windex.setdefault(key, {})
            rec["quality"] = x["category"].lower()
            if x.get("type"): rec["type"] = x["type"]

    talent_of = {}
    for x in season["items"]:
        if x.get("talent"): talent_of[N(x["name"])] = x["talent"]
    for r in mx.get("Gear", []):
        if r.get("Talent") and r["Quality"] in ("Named", "Exotic"): talent_of.setdefault(N(r["Item Name"]), r["Talent"])
    for r in mx.get("Weapons", []):
        if r.get("Talent") and r["Quality"] in ("Named", "Exotic"): talent_of.setdefault(N(r["Name"]), r["Talent"])
    for k, t in mx.get("_wiki_talents", {}).items():
        talent_of.setdefault(N(k), t)
    for collection in (items, windex):
        for name, record in collection.items():
            if record.get("talent"):
                talent_of.setdefault(N(name), record["talent"])

    # ---- Eşyalar (ekipman; uzun liste — içe aktarıcının ana kaynağı)
    eq = []
    for name, v in sorted(items.items(), key=lambda kv: (["exotic", "named", "gear set", "high-end"].index(kv[1].get("quality", "high-end"))
                                                         if kv[1].get("quality") in ("exotic", "named", "gear set", "high-end") else 9,
                                                         kv[1].get("brand") or kv[1].get("gear_set") or "", kv[0])):
        s = season_by.get(N(name), {})
        roll_note = ""
        if v.get("roll_slots"):
            roll_note = "; ".join(f"{k}={val}" for k, val in v["roll_slots"].items() if val not in ("N/A", "type:gear-core", "type:gear-minor"))
        if v.get("mod_slots"):
            roll_note = (roll_note + "; " if roll_note else "") + "Mod yuvaları: " + ", ".join(v["mod_slots"])
        eq.append([name, QUAL.get(v.get("quality"), v.get("quality") or ""), v.get("slot") or "",
                   v.get("brand") or v.get("gear_set") or "", talent_of.get(N(name), v.get("talent") or ""),
                   s.get("season", ""), s.get("desc", "") or roll_note])

    # ---- Markalar
    brands = []
    pieces = defaultdict(lambda: defaultdict(list)); special = defaultdict(list)
    for name, v in items.items():
        b = v.get("brand") or v.get("gear_set")
        if not b or not v.get("slot"): continue
        if v.get("quality") in ("high-end", "gear set"): pieces[b][v["slot"]].append(name)
        else: special[b].append(f"{name} ({SLOT_TR.get(v['slot'], v['slot'])}, {QUAL.get(v.get('quality'), '')})")
    alias = {"Golan Gear Ltd", "Ceska Vyroba"}                    # aynı markanın ikinci yazımı
    sb = season.get("brands", {})
    for b in sorted(d["brands"], key=str.lower):
        if b in alias: continue
        bon = sb.get(b, {}).get("bonuses") or d.get("brand_bonuses", {}).get(b, [])
        bon = (list(bon) + ["", "", ""])[:3]
        brands.append([b] + bon + [", ".join(pieces[b][s]) for s in SLOTS] + ["; ".join(special[b]), sb.get(b, {}).get("season", "")])

    # ---- Gear setler
    gsets = []
    gsb = season.get("gear_sets", {})
    for g in sorted(d["gear_sets"], key=str.lower):
        info = gsb.get(g, {})
        bon = info.get("bonuses") or [f"{i + 2}: {b}" for i, b in enumerate(d.get("gear_set_bonuses", {}).get(g, []))]
        bon = (list(bon) + ["", "", ""])[:3]
        gsets.append([g] + bon + [info.get("chest", ""), info.get("backpack", "")] +
                     [", ".join(pieces[g][s]) for s in SLOTS] + [info.get("season", "")])

    # ---- Exoticler
    exo = []
    ex_names = list(dict.fromkeys(d["exotics"] + [x["name"] for x in season["items"] if x["category"] == "Exotic"] +
                                  [k for k, v in items.items() if v.get("quality") == "exotic"]))
    seen = set()
    for n in ex_names:
        if N(n) in seen: continue
        seen.add(N(n))
        s = season_by.get(N(n), {})
        wt, wdesc = wiki_talent(wiki.get(n))
        wk, ws = wiki_kind(wiki.get(n))
        it = next((v for k, v in items.items() if N(k) == N(n)), None)
        wi = next((v for k, v in windex.items() if N(k) == N(n)), None)
        if s:
            kind, what = s["kind"], s.get("slot") or s.get("type") or ""
        elif it and it.get("slot"):
            kind, what = "gear", it["slot"]
        elif wi:
            kind, what = "weapon", wi.get("type", "")
        elif N(n) in mx_gear:
            kind, what = "gear", mx_gear[N(n)]["Slot"]
        elif N(n) in mx_weap:
            kind, what = "weapon", mx_weap[N(n)]["Weapon Type"]
        else:
            kind, what = wk or "", ws or ""
        tal = s.get("talent") or talent_of.get(N(n)) or wt
        desc = s.get("desc") or wdesc
        exo.append([n, "Ekipman" if kind == "gear" else ("Silah" if kind == "weapon" else ""),
                    SLOT_TR.get(what, what), tal, desc, s.get("season", "")])

    # ---- Named
    named = []
    seen = set()
    for k, v in sorted(items.items()):
        if v.get("quality") != "named": continue
        seen.add(N(k)); s = season_by.get(N(k), {})
        named.append([k, "Ekipman", SLOT_TR.get(v.get("slot"), v.get("slot") or ""), v.get("brand") or "", talent_of.get(N(k), ""), s.get("season", "")])
    for k, v in sorted(windex.items()):
        if v.get("quality") not in ("named", "specialization") or N(k) in seen: continue
        seen.add(N(k)); s = season_by.get(N(k), {})
        base = s.get("base") or (mx_weap.get(N(k), {}) or {}).get("Variant", "")
        named.append([k, "Silah", v.get("type", ""), base, talent_of.get(N(k), ""), s.get("season", "")])
    exo_set = {N(x) for x in ex_names}
    for n in d["named"]:
        if N(n) in seen or N(n) in exo_set: continue
        seen.add(N(n)); s = season_by.get(N(n), {})
        m = mx_weap.get(N(n)) or {}; g = mx_gear.get(N(n)) or {}
        kind = "Silah" if m else ("Ekipman" if g else "")
        named.append([n, kind, m.get("Weapon Type") or SLOT_TR.get(g.get("Slot", ""), g.get("Slot", "")), m.get("Variant") or g.get("Brand", ""),
                      talent_of.get(N(n), ""), s.get("season", "")])

    # ---- Silahlar
    weap = []
    for k, v in sorted(windex.items(), key=lambda kv: (kv[1].get("type", ""), kv[0])):
        m = mx_weap.get(N(k), {})
        bs = v.get("base_stats", {})
        weap.append([k, v.get("type", ""), QUAL.get(v.get("quality"), (v.get("quality") or "").title()),
                     m.get("Variant", "") or v.get("family", ""), _num(m.get("Base Damage")) or bs.get("base_damage"),
                     _num(m.get("RPM")) or bs.get("base_rpm"), _num(m.get("Mag Size")) or bs.get("base_mag_size"),
                     _num(m.get("Reload Speed (ms)")) or bs.get("base_reload_time"), _num(m.get("Optimal Range")) or bs.get("optimal_range"),
                     " ".join(x for x in (m.get("Core 1", ""), f"(max {m['Core 1 Max']})" if m.get("Core 1 Max") else "") if x),
                     " ".join(x for x in (m.get("Core 2", ""), f"(max {m['Core 2 Max']})" if m.get("Core 2 Max") else "") if x),
                     m.get("Attribute 1", ""), talent_of.get(N(k), m.get("Talent", "") if v.get("quality") in ("named", "exotic") else "")])

    # ---- Talentler
    gts = d.get("gear_talent_slots", {})
    mxg = {N(r["Talent"]): r for r in mx.get("GearTalents", [])}
    div2hub_g = {N(r["name"]): r for r in d.get("gear_talent_details", [])}
    st = season.get("talents", {})
    gtal = []
    for t in sorted(set(d["gear_talents"]) | {r["Talent"] for r in mx.get("GearTalents", [])} | set(st.get("gear", {})), key=str.lower):
        r = mxg.get(N(t), {}); info = gts.get(t, {}); dh = div2hub_g.get(N(t), {})
        gtal.append([t, SLOT_TR.get(info.get("slot") or r.get("Slot", ""), info.get("slot") or r.get("Slot", "")),
                     {"normal": "Normal", "perfect": "Perfect", "set": "Set", "exotic": "Exotic"}.get(info.get("kind", ""), info.get("kind", "")),
                     st.get("gear", {}).get(t) or r.get("Desc", "") or dh.get("description", "") or mx.get("_talent_desc", {}).get(t, "")])
    mxw = {N(r["Name"]): r for r in mx.get("WeaponTalents", [])}
    div2hub_w = {N(r["name"]): r for r in d.get("weapon_talent_details", [])}
    cols = ["Assault Rifle", "Rifle", "Marksman Rifle", "SMG", "LMG", "Pistol", "Shotgun"]
    wtal = []
    for t in sorted(set(d["weapon_talents"]) | {r["Name"] for r in mx.get("WeaponTalents", [])} | set(st.get("weapon", {})), key=str.lower):
        r = mxw.get(N(t), {}); dh = div2hub_w.get(N(t), {})
        types = ", ".join(c for c in cols if (r.get(c) or "").strip().lower() == "x") or dh.get("compatibility", "")
        wtal.append([t, types, st.get("weapon", {}).get(t) or r.get("Desc", "") or dh.get("description", "") or mx.get("_talent_desc", {}).get(t, "")])

    # ---- Modlar
    gmods = [[m["stat"], {"offensive": "Saldırı", "defensive": "Savunma", "utility": "Yardımcı"}.get(m.get("type"), m.get("type", "")), m.get("min"), m.get("max"), m.get("fidelity", "")]
             for m in d.get("gear_mods", [])]
    if d.get("weapon_mod_details"):
        wmods = [[r.get("category", ""), r.get("compatibility", ""), r.get("name", ""), r.get("stats", ""), "Div2Hub"]
                 for r in d["weapon_mod_details"]]
        source_mod_names = {N(r.get("name", "")) for r in d["weapon_mod_details"]}
        for name in d.get("weapon_mods", []):
            if N(name) not in source_mod_names:
                wmods.append(["legacy", "", name, "Etki ve uyumluluk doğrulanmamış", "Önceki yerel katalog"])
    else:
        wm_mx = {N(r["Name"]): r for r in mx.get("WeaponMods", [])}
        wmods = []
        for n in d.get("weapon_mods", []):
            r = wm_mx.get(N(n), {})
            wmods.append([r.get("Slot", ""), r.get("Type", ""), n, _num(r.get("valPos")), r.get("pos", ""), _num(r.get("valNeg")), r.get("neg", "")])
        for r in mx.get("WeaponMods", []):
            if N(r["Name"]) not in {N(x) for x in d.get("weapon_mods", [])}:
                wmods.append([r["Slot"], r["Type"], r["Name"], _num(r.get("valPos")), r.get("pos", ""), _num(r.get("valNeg")), r.get("neg", "")])

    # ---- Attribute'lar
    attrs = [[k, ", ".join(a.get("aliases", [])), a.get("max"), a.get("weapon_max"), "Evet" if a.get("core") else "", a.get("unit", "")]
             for k, a in d.get("attributes", {}).items()]

    # ---- Skill'ler
    base_sk = ["Pulse", "Seeker Mine", "Turret", "Sticky Bomb", "Ballistic Shield", "Crusader Shield", "Deflector Shield",
               "Bulwark Shield", "Striker Shield", "Chem Launcher", "Drone", "Hive", "Firefly", "Decoy", "Trap"]
    skills = []
    skill_parents = {N(x["name"]): x["skill"] for x in d.get("skill_variants", [])}
    for s in d.get("skills", []):
        fam = next((b for b in ["Shield", "Hive", "Drone", "Turret", "Seeker", "Firefly", "Trap", "Chem Launcher", "Pulse", "Sticky Bomb", "Decoy"]
                    if b.lower() in s.lower()), "")
        parent = skill_parents.get(N(s), {"Seeker": "Seeker Mine"}.get(fam, fam))
        skills.append([s, parent, "Ana skill" if s in base_sk else "Varyant"])
    skills += [["Riot Foam", "Chem Launcher", "Varyant"]] if not any(x[0] == "Riot Foam" for x in skills) else []

    # ---- Uzmanlıklar
    spec = [[r["Name"], r["Stat"], _num(r["Val"])] for r in mx.get("Specialization", [])]
    sw = d.get("specialization_weapons", {})
    spec_w = [[k, "Tabanca", v] for k, v in sw.get("sidearms", {}).items()] + [[k, "İmza silahı", ""] for k in sw.get("signature", [])]

    # ---- Sezon Y8 ve modifiyeler
    y8 = [[x["season"], x["name"], x["category"], "Ekipman" if x["kind"] == "gear" else "Silah",
           SLOT_TR.get(x.get("slot") or "", x.get("slot") or x.get("type") or ""), x.get("brand") or x.get("base") or "",
           x.get("talent", ""), x.get("desc", "")] for x in season["items"]]
    for g, info in season.get("gear_sets", {}).items():
        y8.append([info["season"], g, "Gear Set", "Ekipman", "6 parça", "", "; ".join(x for x in (info.get("chest"), info.get("backpack")) if x),
                   " · ".join(info["bonuses"])])
    for b, info in season.get("brands", {}).items():
        y8.append([info["season"], b, "Marka", "Ekipman", "6 parça", "", "", " · ".join(info["bonuses"])])
    mods = [[sz, grp, n] for sz, groups in season.get("modifiers", {}).items() for grp, names in groups.items() for n in names]

    return {"eq": eq, "brands": brands, "gsets": gsets, "exo": exo, "named": named, "weap": weap, "gtal": gtal, "wtal": wtal,
            "gmods": gmods, "wmods": wmods, "attrs": attrs, "skills": skills, "spec": spec, "spec_w": spec_w, "y8": y8, "mods": mods,
            "notes": [season.get("brand_bonus_changes", ""), season.get("talent_changes", "")], "version": d.get("version")}


def _num(x):
    if x in (None, ""): return None
    try:
        f = float(str(x).replace(",", "."))
        return int(f) if f.is_integer() else f
    except ValueError:
        return x


# ------------------------------------------------------------------ Excel
FONT = "Arial"
HEAD = PatternFill("solid", fgColor="1F2937")
SHEETS = [
    ("Eşyalar", "eq", ["Eşya adı", "Kalite", "Yuva", "Marka / Set", "Talent", "Sezon", "Not"], [34, 11, 11, 28, 26, 18, 50]),
    ("Markalar", "brands", ["Marka", "1 parça", "2 parça", "3 parça", "Maske", "Sırt çantası", "Göğüslük", "Eldiven", "Holster", "Dizlik",
                            "Named / Exotic parçalar", "Sezon"], [26, 20, 22, 22, 22, 22, 22, 22, 22, 22, 40, 14]),
    ("Gear Setler", "gsets", ["Gear set", "2 parça", "3 parça", "4 parça", "Göğüslük talent", "Sırt talent", "Maske", "Sırt çantası",
                              "Göğüslük", "Eldiven", "Holster", "Dizlik", "Sezon"], [24, 30, 34, 30, 18, 18, 22, 22, 22, 22, 22, 22, 16]),
    ("Exoticler", "exo", ["Exotic", "Tür", "Yuva / Silah türü", "Talent", "Talent açıklaması", "Sezon"], [30, 10, 18, 26, 80, 18]),
    ("Named", "named", ["Named eşya", "Tür", "Yuva / Silah türü", "Marka / Temel silah", "Talent", "Sezon"], [30, 10, 18, 26, 30, 18]),
    ("Silahlar", "weap", ["Silah", "Tür", "Kalite", "Varyant", "Temel hasar", "RPM", "Şarjör", "Doldurma (ms)", "Etkili menzil (m)",
                          "Core 1", "Core 2", "Attribute", "Talent"], [30, 18, 12, 14, 12, 8, 8, 12, 12, 30, 26, 22, 26]),
    ("Ekipman Talentleri", "gtal", ["Talent", "Yuva", "Tür", "Açıklama"], [28, 14, 10, 100]),
    ("Silah Talentleri", "wtal", ["Talent", "Silah türleri", "Açıklama"], [28, 44, 100]),
    ("Ekipman Modları", "gmods", ["Stat", "Tür", "En düşük", "En yüksek", "Adım"], [30, 12, 12, 12, 12]),
    ("Silah Modları", "wmods", ["Kategori", "Uyumluluk", "Mod", "Etkiler", "Kaynak"], [16, 36, 34, 56, 14]),
    ("Attributelar", "attrs", ["Anahtar", "Oyundaki ad(lar)", "En yüksek", "Silahta en yüksek", "Core", "Birim"], [22, 40, 12, 16, 8, 8]),
    ("Skilller", "skills", ["Skill", "Aile", "Tür"], [26, 16, 12]),
    ("Uzmanlıklar", "spec", ["Uzmanlık", "Bonus", "Değer"], [18, 30, 10]),
    ("Uzmanlık Silahları", "spec_w", ["Silah", "Tür", "Uzmanlık"], [32, 14, 16]),
    ("Sezon Y8", "y8", ["Sezon", "Ad", "Kategori", "Tür", "Yuva / Tür", "Marka / Temel", "Talent", "Açıklama / Bonus"], [18, 26, 10, 10, 16, 22, 30, 80]),
    ("Modifiyeler", "mods", ["Sezon", "Grup", "Modifiye"], [18, 20, 36]),
]


def write(data, out):
    wb = Workbook()
    ws = wb.active; ws.title = "Özet"
    for title, key, header, widths in SHEETS:
        sh = wb.create_sheet(title)
        sh.append(header)
        for r in data[key]: sh.append(r)
        for c in sh[1]:
            c.font = Font(name=FONT, bold=True, color="FFFFFF"); c.fill = HEAD
            c.alignment = Alignment(vertical="center", wrap_text=True)
        for row in sh.iter_rows(min_row=2):
            for c in row:
                c.font = Font(name=FONT, size=10)
                if isinstance(c.value, str) and len(c.value) > 60: c.alignment = Alignment(wrap_text=True, vertical="top")
        for i, w in enumerate(widths, 1): sh.column_dimensions[get_column_letter(i)].width = w
        sh.freeze_panes = "B2"
        if sh.max_row > 1:
            ref = f"A1:{get_column_letter(len(header))}{sh.max_row}"
            tname = re.sub(r"[^A-Za-z0-9]", "", title.replace("ı", "i").replace("ş", "s").replace("ç", "c").replace("ü", "u")
                           .replace("ö", "o").replace("ğ", "g").replace("Ö", "O").replace("Ş", "S").replace("Ç", "C").replace("Ü", "U")) or "T"
            tb = Table(displayName="T_" + tname, ref=ref)
            tb.tableStyleInfo = TableStyleInfo(name="TableStyleLight9", showRowStripes=True)
            sh.add_table(tb)

    # Özet: sayılar formülle
    F = lambda **k: Font(name=FONT, **k)
    ws["A1"] = "The Division 2 — eşya kütüphanesi"; ws["A1"].font = F(bold=True, size=15)
    ws["A2"] = f"TD2 Loot Advisor kütüphane sürümü {data['version']} · Son sezon: Y8S3 Red Horizon"; ws["A2"].font = F(italic=True, color="6B7280")
    ws["A4"] = "Bölüm"; ws["B4"] = "Kayıt"; ws["C4"] = "Açıklama"
    for c in ws[4][:3]: c.font = F(bold=True, color="FFFFFF"); c.fill = HEAD
    rows = [
        ("Ekipman (tüm)", "=COUNTA(Eşyalar!A:A)-1", "Eşyalar sayfası: her ekipman parçası, yuvası ve markası/seti"),
        ("  High-End parçalar", '=COUNTIF(Eşyalar!B:B,"High-End")', ""),
        ("  Gear Set parçaları", '=COUNTIF(Eşyalar!B:B,"Gear Set")', ""),
        ("  Named ekipman", '=COUNTIF(Eşyalar!B:B,"Named")', ""),
        ("  Exotic ekipman", '=COUNTIF(Eşyalar!B:B,"Exotic")', ""),
        ("Markalar", "=COUNTA(Markalar!A:A)-1", "Bonuslar ve 6 yuvanın parça adları"),
        ("Gear setler", "=COUNTA('Gear Setler'!A:A)-1", "Bonuslar, göğüslük/sırt talentleri, parça adları"),
        ("Exoticler", "=COUNTA(Exoticler!A:A)-1", "Ekipman ve silah, talent ve açıklama"),
        ("  Exotic silah", '=COUNTIF(Exoticler!B:B,"Silah")', ""),
        ("  Exotic ekipman", '=COUNTIF(Exoticler!B:B,"Ekipman")', ""),
        ("Named eşyalar", "=COUNTA(Named!A:A)-1", "Ekipman ve silah"),
        ("Silahlar", "=COUNTA(Silahlar!A:A)-1", "Tüm silahlar ve varyantları"),
        ("Ekipman talentleri", "=COUNTA('Ekipman Talentleri'!A:A)-1", ""),
        ("Silah talentleri", "=COUNTA('Silah Talentleri'!A:A)-1", ""),
        ("Ekipman modları", "=COUNTA('Ekipman Modları'!A:A)-1", ""),
        ("Silah modları", "=COUNTA('Silah Modları'!A:A)-1", ""),
        ("Attribute'lar", "=COUNTA(Attributelar!A:A)-1", "Programın rulo kalitesini hesapladığı en yüksek değerler"),
        ("Skill'ler", "=COUNTA(Skilller!A:A)-1", ""),
        ("Y8 sezon yenilikleri", "=COUNTA('Sezon Y8'!A:A)-1", "Rise Up, Into the Dark, Red Horizon"),
        ("Sezon modifiyeleri", "=COUNTA(Modifiyeler!A:A)-1", ""),
    ]
    for i, (a, b, c) in enumerate(rows, 5):
        ws.cell(row=i, column=1, value=a).font = F(bold=not a.startswith(" "))
        ws.cell(row=i, column=2, value=b).font = F()
        ws.cell(row=i, column=3, value=c).font = F(color="6B7280")
    r = 5 + len(rows) + 1
    ws.cell(row=r, column=1, value="Red Horizon değişiklikleri").font = F(bold=True)
    for j, n in enumerate([x for x in data["notes"] if x], r + 1):
        ws.cell(row=j, column=1, value=n).font = F(); ws.cell(row=j, column=1).alignment = Alignment(wrap_text=True, vertical="top")
        ws.merge_cells(start_row=j, start_column=1, end_row=j, end_column=3); ws.row_dimensions[j].height = 45
    r = j + 2
    ws.cell(row=r, column=1, value="Programa işleme").font = F(bold=True)
    ws.cell(row=r + 1, column=1, value="Eşyalar, Markalar, Gear Setler, Exoticler, Named, Silahlar, talent ve mod sayfalarındaki düzeltme ve "
                                        "eklemeler 'python scripts/import_library_excel.py' ile programın kütüphanesine (meta/game_data.json) işlenir. "
                                        "Satır silmek kütüphaneden silmez.").font = F()
    ws.merge_cells(start_row=r + 1, start_column=1, end_row=r + 1, end_column=3); ws.row_dimensions[r + 1].height = 45
    ws.cell(row=r + 1, column=1).alignment = Alignment(wrap_text=True, vertical="top")
    ws.cell(row=r + 3, column=1, value="Kaynaklar").font = F(bold=True)
    src = ["Ubisoft sezon notları: Rise Up, Into the Dark, Red Horizon (ubisoft.com/the-division-2/news-updates)",
           "The Division Wiki — thedivision.fandom.com (CC BY-SA 3.0): marka/set parça adları, exotic bilgileri",
           "mx-division-builds — github.com/mxswat/mx-division-builds (CC BY-NC-SA 4.0): talent açıklamaları, modlar, silah değerleri",
           "keengamer.com, mp1st.com, skycoach.gg: sezon ayrıntıları ve silah türleri",
           "Boş hücre: kaynaklarda bilgi yok (ör. Teapot/Steamer silah türü, Tip of the Spear parça adları)."]
    for j, s in enumerate(src, r + 4): ws.cell(row=j, column=1, value=s).font = F(size=9, color="374151")
    ws.column_dimensions["A"].width = 34; ws.column_dimensions["B"].width = 10; ws.column_dimensions["C"].width = 70
    wb.save(out)
    return out


if __name__ == "__main__":
    out = Path(sys.argv[1]) if len(sys.argv) > 1 else ROOT / "meta" / "TD2_Kutuphane.xlsx"
    data = build(*load())
    write(data, out)
    print(out, {k: len(v) for k, v in data.items() if isinstance(v, list)})
